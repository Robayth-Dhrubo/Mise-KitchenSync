"""Extract a Faberge Library download into a reviewable, lossless import package.

Original bytes and source grids are retained separately from editable records.
Run: python3 scripts/import-faberge.py SOURCE_DIRECTORY OUTPUT_DIRECTORY
No credentials or customer data are written into the source checkout.
"""
from __future__ import annotations
import csv, hashlib, io, json, re, sys
from collections import defaultdict
from datetime import date, datetime, time
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

root, out = map(Path, sys.argv[1:3])
out.mkdir(parents=True, exist_ok=True)
(out / 'assets').mkdir(exist_ok=True)
sources = json.loads((root / 'source-manifest.json').read_text())
assets, files, source_by_name = {}, [], {}

def digest(raw): return hashlib.sha256(raw).hexdigest()
def put(raw, local=None):
    key = digest(raw)
    if key not in assets:
        p = local or out / 'assets' / key
        if not local: p.write_bytes(raw)
        assets[key] = {'sha256': key, 'path': str(p), 'bytes': len(raw)}
    return key
def safe(v):
    if isinstance(v, (datetime, date, time)): return v.isoformat()
    if isinstance(v, float) and (v != v or abs(v) == float('inf')): return None
    return v
def clean(v, length=300): return str(v or '').strip()[:length]
def number(v):
    if isinstance(v, bool): return None
    if isinstance(v, (int, float)): return v if v == v else None
    try: return float(str(v).replace(',', '').replace('$', '').strip())
    except (ValueError, TypeError): return None
def sid(prefix, text): return prefix + '-' + digest(text.encode())[:16]
def rowdata(index, values, formulas=None, formats=None):
    values = [safe(v) for v in values]
    while values and values[-1] is None: values.pop()
    result = {'row': index, 'cells': values}
    if formulas: result['formulas'] = formulas
    if formats: result['formats'] = formats
    return result
def csv_rows(p):
    raw = p.read_text(encoding='utf-8-sig')
    if raw.startswith('sep='):
        first, raw = raw.split('\n', 1); delimiter = first.strip()[4:5]
    else:
        try: delimiter = csv.Sniffer().sniff(raw[:8192]).delimiter
        except csv.Error: delimiter = ','
    return list(csv.reader(io.StringIO(raw), delimiter=delimiter))
def table_asset(file, name, rows, columns):
    raw = json.dumps({'name': name, 'columns': columns, 'rows': rows}, ensure_ascii=False, separators=(',', ':')).encode()
    key = put(raw)
    table = {'id': sid('table', file['id'] + name), 'name': name, 'asset': key, 'rows': len(rows), 'columns': len(columns)}
    file['tables'].append(table)

for source in sources:
    p = root / source['name']
    if not p.is_file(): raise RuntimeError('Missing source: ' + p.name)
    key = put(p.read_bytes(), p)
    file = {'id': sid('source', source['library_file_id']), 'name': p.name, 'asset': key,
            'bytes': p.stat().st_size, 'modifiedAt': source['modified_at'], 'tables': [], 'category': 'Operations'}
    lname = p.name.lower()
    for match, category in [('price', 'Prices'), ('costing', 'Recipes'), ('recipe', 'Recipes'), ('payroll', 'Payroll'),
                            ('schedule', 'Staffing'), ('supplier', 'Suppliers'), ('sales', 'Sales'), ('2026-09-01', 'Sales'),
                            ('engineering', 'Menu analysis'), ('purchase history', 'Purchasing'), ('menu.pdf', 'Menu')]:
        if match in lname: file['category'] = category; break
    file['status'] = 'reference' if any(x in lname for x in ['results(2)', 'updated_hollandaise', 'executive_chef', 'food costing']) else 'current'
    if any(x in lname for x in ['recipe_bible', 'schedule', 'prep_system']): file['status'] = 'draft'
    files.append(file); source_by_name[p.name] = file
    if p.suffix.lower() == '.xlsx':
        cached = load_workbook(p, data_only=True)
        formula_book = load_workbook(p, data_only=False)
        for sheet in formula_book:
            values_sheet = cached[sheet.title]
            rows, redacted = [], set()
            for rr in sheet:
                # Password fields remain in the private original, not in table previews.
                for c in rr:
                    if c.row < 12 and isinstance(c.value, str) and re.search(r'password|mot de passe|vault ref', c.value, re.I): redacted.add(c.column - 1)
                values, formulas, formats = [], {}, {}
                for c in rr:
                    idx = c.column - 1; value = values_sheet.cell(c.row, c.column).value
                    if idx in redacted and c.row > 4 and c.value is not None: value = '[See private original]'
                    elif c.data_type == 'f': formulas[str(idx)] = c.value
                    if value is not None and c.number_format != 'General': formats[str(idx)] = c.number_format
                    values.append(value)
                if any(v is not None for v in values) or formulas: rows.append(rowdata(rr[0].row, values, formulas, formats))
            table_asset(file, sheet.title, rows, [get_column_letter(i) for i in range(1, sheet.max_column + 1)])
    elif p.suffix.lower() == '.csv':
        rows = csv_rows(p)
        table_asset(file, 'CSV records', [rowdata(i, row) for i, row in enumerate(rows[1:], 2)], rows[0])
    elif p.suffix.lower() == '.docx':
        document = Document(p); rows = []
        for block in document.element.body:
            if block.tag.endswith('}p'):
                text = Paragraph(block, document).text
                if text: rows.append(rowdata(len(rows) + 1, [text]))
            elif block.tag.endswith('}tbl'):
                for rr in Table(block, document).rows:
                    cells = [c.text for c in rr.cells]
                    if any(cells): rows.append(rowdata(len(rows) + 1, cells))
        width = max((len(rr['cells']) for rr in rows), default=1)
        table_asset(file, 'Recipe document', rows, ['Content'] + ['Column ' + str(i) for i in range(2, width + 1)])
    elif p.suffix.lower() in ['.txt', '.md']:
        table_asset(file, 'Notes', [rowdata(i, [line]) for i, line in enumerate(p.read_text().splitlines(), 1) if line], ['Content'])

def provenance(file, sheet, row=None, status='reference', detail='', price_date=None):
    result = {'fileId': source_by_name[file]['id'], 'file': file, 'sheet': sheet, 'status': status, 'detail': clean(detail, 1000)}
    if row is not None: result['row'] = row
    if price_date: result['priceDate'] = safe(price_date)[:10]
    return result

workspace = {'version': 1, 'sampleData': False, 'settings': {'kitchenName': 'Fabergé', 'targetFoodCost': 30},
             'ingredients': [], 'recipes': [], 'suppliers': [], 'prep': [], 'waste': [], 'orders': []}
suppliers = {}
def supplier(name):
    if name not in suppliers:
        suppliers[name] = {'id': sid('fab-supplier', name), 'name': name, 'contact': '', 'email': '', 'phone': '', 'deliveryDays': '', 'notes': ''}
    return suppliers[name]

master_name = 'Faberge_Master_Price_List_Corrected(1).xlsx'
sheet = load_workbook(root / master_name, data_only=True)['Master item list']
names = defaultdict(int)
for row in list(sheet.values)[7:]:
    if row[2]: names[str(row[2])] += 1
products = {}
price_issues = []
for index, rr in enumerate(list(sheet.values)[7:], 8):
    if not rr[2]: continue
    category, vendor, product, brand, code = map(clean, rr[:5])
    name = product if names[product] == 1 else f'{product} · {brand or code}'
    raw_qty, raw_unit = number(rr[12]), rr[13]
    unit, qty = 'each', None
    if raw_unit in ['kg', 'g']: unit, qty = 'g', raw_qty * (1000 if raw_unit == 'kg' else 1) if raw_qty is not None else None
    elif raw_unit in ['L', 'ml']: unit, qty = 'ml', raw_qty * (1000 if raw_unit == 'L' else 1) if raw_qty is not None else None
    elif raw_unit == 'each': unit, qty = 'each', raw_qty
    elif raw_unit == 'oz' and category in ['Baking & Pastry', 'Pantry & Condiments', 'Bread & Bakery']:
        unit, qty = 'g', raw_qty * 28.349523125 if raw_qty is not None else None
    detail = '; '.join(str(v) for v in [rr[18], rr[19], rr[21]] if v)
    flag = clean(rr[19])
    status = 'estimate' if re.search('ESTIMATE|RESEARCH', detail, re.I) else 'confirmed'
    if qty is None:
        status = 'review'; detail += f'; Confirm quantity and unit from invoice ({raw_qty} {raw_unit}).'
    if product in ['Ham Boneless Toupie', 'Salmon Atlantic Smoked Sliced']:
        qty = None; status = 'review'
        detail += '; Price basis disputed by the 8 October Benedict costing. Confirm per kg/per lb versus full pack.'
    if flag and status == 'confirmed': status = 'review'
    if status != 'confirmed': price_issues.append({'product': name, 'issue': detail})
    sup = supplier(vendor)
    contact = clean(rr[20]); email = re.search(r'[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}', contact)
    phone = re.search(r'\d{3}[- ]\d{3}[- ]\d{4}', contact)
    if email and not sup['email']: sup['email'] = email.group()
    if phone and not sup['phone']: sup['phone'] = phone.group()
    item = {'id': sid('fab-product', vendor + '|' + product + '|' + brand + '|' + code), 'name': name[:120], 'category': category,
            'supplierId': sup['id'], 'purchaseLabel': clean(rr[5]), 'baseUnit': unit, 'packQuantity': qty,
            'price': number(rr[10]), 'yieldPct': 100, 'stock': None, 'par': None,
            'notes': clean((flag + '; ' if flag else '') + 'As-purchased quantity; usable yield unmeasured. ' + clean(rr[21])),
            'source': provenance(master_name, sheet.title, index, status, detail, rr[17])}
    workspace['ingredients'].append(item); products[product] = item

# The supplier directory preserves spelling from the current price list. The
# separately named contacts stay separate unless their identities are explicit.
contact_file = 'Faberge_Supplier_Contacts_and_Logins.xlsx'
contact_sheet = load_workbook(root / contact_file, data_only=True).worksheets[0]
for index, rr in enumerate(list(contact_sheet.values)[4:15], 5):
    if not rr[0]: continue
    name = str(rr[0]); name = name + ' · ' + str(rr[1]) if name == 'Bread Supplier' else name
    if name == 'Aliments Viens': name = 'Aliments Phillip Viens'
    sup = supplier(name)
    sup['contact'] = clean(rr[1]); sup['phone'] = clean(' / '.join(str(v) for v in rr[3:6] if v))
    if rr[6]: sup['email'] = clean(rr[6], 150)
    sup['notes'] = clean('; '.join(str(v) for v in [rr[2], rr[7]] if v))
    sup['source'] = provenance(contact_file, contact_sheet.title, index, 'reference')
workspace['suppliers'] = list(suppliers.values())

# Current working recipes retain missing portions, prices, and ingredient
# quantities. No historical selling price is substituted for a new dish.
bible_file = 'Faberge_New_Menu_Recipe_Bible_Working_Draft.docx'
bible = Document(root / bible_file)
paragraphs = [p.text.strip() for p in bible.paragraphs if p.text.strip()]
category = 'Working menu'
recipe_by_name = {}
for index, text in enumerate(paragraphs):
    if text.startswith(('BRUNCH —', 'SALADS &', 'THE CLASSICS', 'LUNCH —', 'KIDS —', 'COMPONENT RECIPE')): category = text
    if text != 'INGREDIENTS' or index < 2: continue
    if paragraphs[index - 2] == 'COMPONENT RECIPE': name, description = paragraphs[index - 1], 'Component recipe: quantities and yield not yet supplied.'
    else: name, description = paragraphs[index - 2], paragraphs[index - 1]
    if len(name) > 120 or name == 'INGREDIENTS': continue
    item = {'id': sid('fab-recipe', name), 'name': name, 'category': category, 'portions': None, 'sellingPrice': None,
            'notes': clean(description), 'items': [], 'source': provenance(bible_file, 'Recipe document', index + 1, 'draft', 'Working draft; approval, quantities and yield blank in source.')}
    recipe_by_name[name] = item

prep_file = 'Faberge_Operating_Control_Center_2026.xlsx'
prep_sheet = load_workbook(root / prep_file, data_only=True)['PREP BOARD']
for index, rr in enumerate(list(prep_sheet.values)[5:24], 6):
    if not rr[0]: continue
    workspace['prep'].append({'id': sid('fab-prep', str(rr[0])), 'name': clean(rr[0], 120), 'station': clean(rr[1]), 'assignee': clean(rr[7]),
                              'quantity': '', 'due': None, 'priority': 'normal', 'status': 'todo',
                              'source': provenance(prep_file, prep_sheet.title, index, 'draft', 'Undated prep template; targets, stock, assignment and priority are blank.')})

unpriced = {}
aliases = {'Pasteurized egg yolks': 'Egg Yolk Frozen Pasteurized', 'Canola oil': 'Oil Canola Jib', 'Ground black pepper': 'Pepper Black Ground 32M',
           'Granulated sugar': 'Sugar Granulated Fine Paper', 'White sugar': 'Sugar Granulated Fine Paper', 'Fine white sugar, for the caramelized surface': 'Sugar Granulated Fine Paper',
           'Onion powder or finely ground onion granules': 'Spice Onion Granulated', 'Peeled garlic, lightly crushed': 'Peeled garlic',
           'Peeled garlic, minced': 'Peeled garlic', 'Café mushrooms, raw purchase weight': 'Café mushrooms', 'Baby spinach, raw purchase weight': 'Baby spinach',
           'Tomatoes 6x6, raw purchase weight': 'Tomatoes, 6x6', 'Whole potatoes, washed': 'Potatoes', 'All-purpose flour': 'Flour Bakersource All Purpose',
           'Light brown sugar': 'Sugar Brown Golden Yellow'}
def ingredient_for(name, unit, file):
    base = 'ml' if unit == 'ml' else 'each' if unit == 'each' else 'g'
    product = products.get(aliases.get(name, name))
    if product and product['baseUnit'] == base: return product
    key = name + '|' + base
    if key not in unpriced:
        unpriced[key] = {'id': sid('fab-recipe-input', key), 'name': clean(name + (' (weighed)' if base == 'g' and name in ['35% cream', 'Cream, 35%', 'Whole milk, 3.25%'] else ''), 120),
                         'category': 'Recipe inputs to confirm', 'supplierId': None, 'purchaseLabel': '', 'baseUnit': base, 'packQuantity': None, 'price': None,
                         'yieldPct': 100, 'stock': None, 'par': None, 'notes': 'Recipe input. Confirm purchasing specification and invoice price; no density or yield assumed.',
                         'source': provenance(file, 'Recipe document', None, 'review', 'No confirmed compatible purchasing specification was supplied for this recipe input.')}
    return unpriced[key]
def batch_recipe(file, table_index, name, portions, detail):
    table = Document(root / file).tables[table_index]
    item = {'id': sid('fab-recipe', name), 'name': name, 'category': 'Trial batch recipes', 'portions': portions, 'sellingPrice': None, 'notes': clean(detail),
            'items': [], 'source': provenance(file, 'Recipe document', None, 'draft', detail)}
    for index, row in enumerate(table.rows[1:], 1):
        ingredient_name, amount = [c.text for c in row.cells[:2]]
        match = re.search(r'([\d,.]+)\s*(g|mL|ml|pod)', amount)
        if not match: continue
        quantity = float(match.group(1).replace(',', '')); unit = 'each' if match.group(2) == 'pod' else 'ml' if match.group(2).lower() == 'ml' else 'g'
        ingredient = ingredient_for(ingredient_name, unit, file)
        item['items'].append({'id': 'line-' + str(index), 'ingredientId': ingredient['id'], 'quantity': quantity, 'unit': unit})
    recipe_by_name[name] = item

benedict = 'Faberge_Eggs_Benedict_Recipe.docx'
batch_recipe(benedict, 1, 'Hollandaise', 530 / 60, '8 October trial: about 530 g usable at 60 g/plate. Yield unmeasured; full recipe and method in source. Confirm butter, MSG, cayenne and siphon gas costs.')
gas = ingredient_for('Siphon gas', 'each', benedict)
recipe_by_name['Hollandaise']['items'].append({'id': 'line-gas', 'ingredientId': gas['id'], 'quantity': None, 'unit': 'each'})
batch_recipe(benedict, 2, 'Breakfast Potatoes', 1700 / 150, 'Trial yield 1.70 kg; 150 g per plate. Weigh finished batch before relying on portion costing.')
batch_recipe(benedict, 3, 'Confit Tomatoes', 700 / 60, 'Trial yield 700 g cooked/drained; 60 g per plate. Yield requires measurement.')
batch_recipe(benedict, 4, 'Mushroom & spinach filling', 10, 'Trial yield 800 g; 80 g per plate. Yield requires measurement.')
pain = 'Faberge_Pain_Perdu_Brioche_Recipe(1).docx'
for table_index, name, portions, detail in [(0, 'Pain perdu · brioche and custard', 6, 'Trial base for 6 plates; custard absorption and usable yield unmeasured. Cooking butter, fruit and garnishes are separate.'),
                                          (2, 'Orange-vanilla mascarpone', 270 / 40, 'Trial yield about 270 g before bowl loss; 40 g per plate.'),
                                          (3, 'Lightly sweetened oat crumble', 230 / 8, 'Trial baked yield about 230 g; 8 g per plate.'),
                                          (4, 'Vanilla candied pecans', 600 / 8, 'Trial baked yield about 600 g; 8 g per plate.'),
                                          (5, 'Caramelized pears', 6, 'Trial fruit batch for 6 plates; target 70 g cooked fruit per plate. Yield unmeasured.')]:
    batch_recipe(pain, table_index, name, portions, detail)
recipe_by_name['Pain perdu brioché'] = {'id': sid('fab-recipe', 'Pain perdu brioché'), 'name': 'Pain perdu brioché', 'category': 'Sweet breakfast', 'portions': None, 'sellingPrice': None, 'items': [],
                                     'notes': 'Latest 9 October trial recipe. Use the original recipe document for the complete assembly, components, cooking and plating; approved yield and selling price remain blank.',
                                     'source': provenance(pain, 'Recipe document', None, 'draft', 'Replaces the earlier generic brioche-toast description for this development trial.')}
workspace['ingredients'].extend(unpriced.values())
workspace['recipes'] = list(recipe_by_name.values())

# POS item detail is the daily/item fact table. Other report totals are controls,
# not extra revenue. Keep options separate because their daily allocation is absent.
pos_name = next(p.name for p in root.glob('By Items,*.csv'))
pos_rows = csv_rows(root / pos_name)
records = [dict(zip(pos_rows[0], rr)) for rr in pos_rows[1:] if len(rr) == len(pos_rows[0]) and re.match(r'^2026-\d{2}-\d{2}', rr[2])]
daily, item_sales, category_sales = {}, {}, {}
receipts = defaultdict(set)
for rr in records:
    day = rr['Date'][:10]; qty, total = number(rr['Quantity']), number(rr['Total'])
    if qty is None or total is None: raise RuntimeError('Unparseable POS numeric row')
    receipts[day].add(rr['Receipt #'])
    for table, key in [(daily, day), (item_sales, (rr['Item ID'], rr['Name'], rr['Category'])), (category_sales, rr['Category'])]:
        if key not in table: table[key] = {'quantity': 0, 'sales': 0}
        table[key]['quantity'] += qty; table[key]['sales'] += total
sales_daily = [{'date': day, 'sales': round(value['sales'], 2), 'quantity': value['quantity'], 'receipts': len(receipts[day])} for day, value in sorted(daily.items())]
sales_items = [{'id': key[0], 'name': key[1], 'category': key[2], 'quantity': value['quantity'], 'sales': round(value['sales'], 2)} for key, value in item_sales.items()]
sales_categories = [{'name': key, 'quantity': value['quantity'], 'sales': round(value['sales'], 2)} for key, value in category_sales.items()]
controls = []
for prefix in ['By Items', 'Net sales by category', 'Categories', 'Account', 'Extra - Options']:
    p = next(root.glob(prefix + ',*.csv')); raw = p.read_text(); match = re.search(r'Grand-total:\s*([\d,.]+)', raw)
    controls.append({'report': prefix, 'total': number(match.group(1)) if match else None, 'fileId': source_by_name[p.name]['id']})
hour_file = next(root.glob('Hourly sales,*.csv')); hour_rows = csv_rows(hour_file)
hours = [{'hour': label, 'sales': number(value)} for label, value in zip(hour_rows[0][1:-1], hour_rows[1][1:-1])]

engineering_file = 'Faberge_Menu_Engineering_EXACT_FORMAT_Sep_Oct_2026.xlsx'
menu_sheet = load_workbook(root / engineering_file, data_only=True)['Menu Engineering']
menu = []
for rr in list(menu_sheet.values)[4:]:
    if isinstance(rr[0], str) and isinstance(rr[2], (int, float)) and rr[0] not in ['TOTAL', 'COLUMN TOTALS']:
        menu.append({'name': rr[0], 'category': rr[1] or '', 'sold': rr[2], 'foodCost': number(rr[4]), 'sellingPrice': number(rr[5]), 'revenue': number(rr[9]),
                     'classification': rr[13] or 'No sales', 'basis': rr[14] or 'Historical cost'})

schedule_name = 'faberge_kitchen_schedule.xlsx'
schedule_sheet = load_workbook(root / schedule_name, data_only=True).worksheets[0]
schedule = []
for rr in list(schedule_sheet.values)[11:17]:
    salary = rr[2] == 'Salary'
    schedule.append({'name': rr[0], 'role': rr[1], 'payType': rr[2], 'rate': number(rr[3]),
                     'days': [number(v) for v in rr[4:11]], 'hours': None if salary else number(rr[11]), 'weeklyCost': number(rr[12]),
                     'dayUnit': 'work day' if salary else 'hours'})
payroll_file = 'Faberge Payroll Tracking .xlsx'
payroll_book = load_workbook(root / payroll_file, data_only=True)
payroll = []
for sheet in payroll_book:
    for rr in sheet.values:
        if len(rr) > 19 and isinstance(rr[1], str) and re.match(r'^\d{2}/\d{2}/\d{2}-\d{2}/\d{2}/\d{2}', rr[1]):
            payroll.append({'year': sheet.title, 'period': rr[1], 'sales': number(rr[2]), 'foh': number(rr[3]), 'boh': number(rr[7]), 'admin': number(rr[11]),
                            'management': number(rr[15]), 'reportedTotal': number(rr[19])})
current = payroll_book['2026-2027 Payroll']
payroll_summary = {'year': current.title, 'netSales': number(current['B2'].value), 'foh': number(current['B3'].value), 'boh': number(current['B4'].value),
                   'admin': number(current['B5'].value), 'management': number(current['B6'].value), 'reportedTotal': number(current['B7'].value),
                   'coverage': '6 April–20 September 2026', 'fileId': source_by_name[payroll_file]['id']}
summary = {'salesRows': len(records), 'salesTotal': round(sum(number(rr['Total']) for rr in records), 2), 'salesFrom': min(daily), 'salesThrough': max(daily),
           'salesDates': len(daily), 'receipts': sum(len(v) for v in receipts.values()), 'stockCounts': 0, 'priceProducts': 174, 'recipeInputs': len(unpriced),
           'recipes': len(workspace['recipes']), 'suppliers': len(suppliers), 'prepTemplates': len(workspace['prep']), 'files': len(files), 'tables': sum(len(f['tables']) for f in files)}
report = {'batchId': 'faberge-2026-10-10', 'title': 'Fabergé data import', 'importedAt': datetime.now().isoformat() + 'Z', 'files': files, 'summary': summary,
          'sales': {'daily': sales_daily, 'items': sorted(sales_items, key=lambda r: -r['sales']), 'categories': sorted(sales_categories, key=lambda r: -r['sales']),
                    'controls': controls, 'hours': hours, 'fileId': source_by_name[pos_name]['id'], 'menu': menu},
          'team': {'schedule': schedule, 'scheduleFileId': source_by_name[schedule_name]['id'], 'payroll': payroll, 'payrollSummary': payroll_summary},
          'issues': [{'title': 'Inventory counts and pars are blank', 'detail': '174 price-list products imported. No current physical stock counts were supplied; stock and pars remain unknown.'},
                     {'title': 'Recipe drafts need approval', 'detail': 'The working bible has blank quantities and yields. Benedict and Pain Perdu batch recipes retain their trial yields; original preparation methods are available in Sources.'},
                     {'title': 'POS totals differ', 'detail': 'Item detail totals $98,541.90. Account totals $99,326.62 and options $784.67. Adding options to item detail leaves $0.05 unreconciled. Category and hourly exports have other control totals; they are kept separately.'},
                     {'title': 'Roster is an undated draft', 'detail': 'The $4,210 weekly kitchen plan is not approved actual payroll. The chef row records work days, not hours.'},
                     {'title': 'Payroll management column needs reconciliation', 'detail': 'Reported current-year payroll excludes the separately shown $13,819.29 Management column. It has not been added to payroll totals.'},
                     {'title': 'Price basis needs confirmation', 'detail': 'Ham and smoked salmon rates have conflicting per-weight/full-pack descriptions. Their converted unit costs are blocked until confirmed; other estimated packs keep their source flags.'}],
          'priceIssues': price_issues}
payload = {'workspace': workspace, 'report': report}
(out / 'import.json').write_text(json.dumps(payload, ensure_ascii=False, separators=(',', ':')))
(out / 'assets.json').write_text(json.dumps(list(assets.values()), separators=(',', ':')))
print(json.dumps({'summary': summary, 'payloadBytes': (out / 'import.json').stat().st_size, 'assets': len(assets), 'assetBytes': sum(v['bytes'] for v in assets.values())}))
